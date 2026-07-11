"""DOCX document parser."""

from pathlib import Path

from docx import Document

from app.core.enums import KnowledgeSourceType
from app.knowledge.parser.base import DocumentParser, ParsedDocument, ParsedPage


class DocxDocumentParser(DocumentParser):
    @property
    def supported_types(self) -> frozenset[KnowledgeSourceType]:
        return frozenset({KnowledgeSourceType.DOCX})

    async def parse_file(self, file_path: str, *, source_type: KnowledgeSourceType) -> ParsedDocument:
        path = Path(file_path)
        document = Document(str(path))
        paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        text = "\n\n".join(paragraphs)
        headings = [p.text for p in document.paragraphs if p.style and p.style.name.startswith("Heading")]
        return ParsedDocument(
            text=text,
            pages=[ParsedPage(page_number=1, content=text, metadata={"headings": headings[:20]})],
            metadata={"filename": path.name, "source_type": source_type.value, "paragraph_count": len(paragraphs)},
        )

    async def parse_url(self, url: str) -> ParsedDocument:
        raise NotImplementedError("DOCX URL parsing is not supported")
